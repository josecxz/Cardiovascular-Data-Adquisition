"""
Get the version-1 posting date for every Research Square preprint in the workbook
whose recorded DOI is not /v1.

Research Square registers a separate DOI per version (.../v1, .../v2, ...), so the
v1 date is obtainable three ways, tried in this order per DOI:

    1. crossref  - api.crossref.org/works/<v1 doi>   ("posted" date)
    2. web       - the researchsquare.com article page embeds a __NEXT_DATA__ JSON
                   blob whose "nonDraftVersions" lists every version with its date
    3. selenium  - the same page rendered in a real Chrome browser (last resort,
                   only used for DOIs the first two tiers could not resolve)

Results are cached in a JSON file, so re-runs are resumable and cost no requests.

Usage:
    python Get_V1_Publish_Dates.py                  # tiers 1 + 2
    python Get_V1_Publish_Dates.py --selenium       # also allow the browser fallback
    python Get_V1_Publish_Dates.py --limit 20       # smoke test
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import date as _date

import openpyxl
import requests

DEFAULT_INPUT = 'medRxiv and RS CV preprints data analysis 06_30.xlsx'
DEFAULT_SHEET = 'Final analytic dataset'
CACHE_FILE = 'v1_dates_cache.json'
# Crossref asks callers to identify themselves ("polite pool") and gives identified
# traffic better service. Set CROSSREF_CONTACT to an email to opt in; without it the
# scripts still work, they just run in the anonymous pool.
CONTACT = os.environ.get('CROSSREF_CONTACT', '')
USER_AGENT = (f'preprint-v1-date-collector/1.0 (mailto:{CONTACT})' if CONTACT
              else 'preprint-v1-date-collector/1.0')

VERSION_RE = re.compile(r'/v(\d+)\s*$')
NEXT_DATA_RE = re.compile(
    r'<script id="__NEXT_DATA__" type="application/json">(.*?)</script>', re.S)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def version_of(doi):
    """Return the version number encoded in a DOI, or None if it has none."""
    m = VERSION_RE.search(doi or '')
    return int(m.group(1)) if m else None


def to_v1(doi):
    return VERSION_RE.sub('/v1', doi.strip())


def rs_article_url(doi, version='v1', sheet_url=None):
    """Article URL for a given version.

    The URL recorded in the sheet is authoritative and is simply re-pointed at the
    wanted version. Deriving it from the DOI only works for the modern rs.3.rs-NNN
    format -- the older rs.2.NNNNN DOIs use an unrelated rs-NNNN slug, so without
    the sheet's URL those cannot be reached on the web.
    """
    if sheet_url:
        u = str(sheet_url).strip()
        if u.startswith('http'):
            return VERSION_RE.sub('/' + version, u)
    if '/' not in doi:
        return None
    stem = doi.rsplit('/', 1)[0]                      # 10.21203/rs.3.rs-28577
    if '/' not in stem:
        return None
    slug = stem.split('/', 1)[1]                      # rs.3.rs-28577  or  rs.2.10364
    if 'rs.3.' not in slug:
        return None                                   # old format: not derivable
    return f'https://www.researchsquare.com/article/{slug.split("rs.3.")[-1]}/{version}'


def parse_loose_date(value):
    """Normalise the page's human dates ("May 19th, 2021") to 'YYYY-MM-DD'."""
    text = str(value or '').strip()
    if not text:
        return None
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}.*', text):
        return text[:10]
    text = re.sub(r'(\d+)(st|nd|rd|th)', r'\1', text)      # 19th -> 19
    for fmt in ('%B %d, %Y', '%b %d, %Y', '%d %B %Y', '%d %b %Y', '%B %d %Y'):
        try:
            return time.strftime('%Y-%m-%d', time.strptime(text, fmt))
        except ValueError:
            pass
    return None


def iso(date_parts):
    """Crossref [[Y, M, D]] -> 'YYYY-MM-DD' (month/day may be missing)."""
    if not date_parts:
        return None
    p = list(date_parts[0]) + [1, 1]
    return f'{p[0]:04d}-{p[1]:02d}-{p[2]:02d}' if p[0] else None


def days_between(a, b):
    try:
        ya, ma, da = (int(x) for x in str(a)[:10].split('-'))
        yb, mb, db = (int(x) for x in str(b)[:10].split('-'))
        return (_date(yb, mb, db) - _date(ya, ma, da)).days
    except Exception:
        return None


def make_session():
    s = requests.Session()
    s.headers['User-Agent'] = USER_AGENT
    return s


# --------------------------------------------------------------------------- #
# tier 1 - Crossref
# --------------------------------------------------------------------------- #
def date_from_crossref(session, v1_doi):
    r = session.get(f'https://api.crossref.org/works/{v1_doi}', timeout=30)
    if r.status_code == 404:
        return None, 'not in crossref'
    r.raise_for_status()
    msg = r.json()['message']
    for field in ('posted', 'published', 'issued', 'created'):
        d = iso(msg.get(field, {}).get('date-parts'))
        if d:
            return d, field
    return None, 'no date field'


# --------------------------------------------------------------------------- #
# tier 2 - the article page's embedded JSON
# --------------------------------------------------------------------------- #
def versions_from_next_data(html):
    """-> {version_number: 'YYYY-MM-DD'} parsed out of the page's __NEXT_DATA__."""
    m = NEXT_DATA_RE.search(html)
    if not m:
        return {}
    try:
        init = json.loads(m.group(1))['props']['pageProps']['initialData']
    except Exception:
        return {}
    out = {}
    for v in init.get('nonDraftVersions') or []:
        n = version_of(v.get('doi', ''))
        d = parse_loose_date(v.get('date'))
        if n and d:
            out[n] = d
    # A superseded v1 is moved out of nonDraftVersions into archivedVersions, where
    # it carries no DOI. On the v1 page itself, postedDate is that missing v1 date.
    if 1 not in out and version_of(init.get('doi', '')) == 1:
        posted = parse_loose_date(init.get('postedDate'))
        if not posted:
            archived = [parse_loose_date(a.get('date'))
                        for a in (init.get('archivedVersions') or [])]
            archived = sorted(d for d in archived if d)
            posted = archived[0] if archived else None
        if posted:
            out[1] = posted
    return out


def date_from_web(session, doi, sheet_url=None):
    url = rs_article_url(doi, 'v1', sheet_url)
    if not url:
        return None, 'no url', {}
    r = session.get(url, timeout=45)
    if r.status_code != 200:
        return None, f'http {r.status_code}', {}
    versions = versions_from_next_data(r.text)
    return versions.get(1), ('ok' if versions.get(1) else 'no v1 in page'), versions


# --------------------------------------------------------------------------- #
# tier 3 - Selenium / Chrome
# --------------------------------------------------------------------------- #
class Browser:
    """Lazily-started headless Chrome; costs nothing unless actually used."""

    def __init__(self):
        self.driver = None

    def _start(self):
        from selenium import webdriver
        from selenium.webdriver.chrome.options import Options
        opts = Options()
        opts.add_argument('--headless=new')
        opts.add_argument('--disable-gpu')
        opts.add_argument('--no-sandbox')
        opts.add_argument('--window-size=1280,900')
        opts.add_argument(f'--user-agent={USER_AGENT}')
        # Selenium 4 downloads a matching chromedriver by itself.
        self.driver = webdriver.Chrome(options=opts)
        self.driver.set_page_load_timeout(60)

    def date_for(self, doi, sheet_url=None):
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support import expected_conditions as EC
        from selenium.webdriver.support.ui import WebDriverWait

        url = rs_article_url(doi, 'v1', sheet_url)
        if not url:
            return None, 'no url', {}
        if self.driver is None:
            self._start()
        self.driver.get(url)
        WebDriverWait(self.driver, 30).until(
            EC.presence_of_element_located((By.ID, '__NEXT_DATA__')))
        blob = self.driver.execute_script(
            "return document.getElementById('__NEXT_DATA__').textContent;")
        versions = versions_from_next_data(
            f'<script id="__NEXT_DATA__" type="application/json">{blob}</script>')
        if versions.get(1):
            return versions[1], 'ok', versions
        # Fall back to the human-readable "Posted DD Month YYYY" line.
        text = self.driver.find_element(By.TAG_NAME, 'body').text
        m = re.search(r'Posted\s+(\d{1,2}\s+\w+,?\s+\d{4})', text)
        if m:
            for fmt in ('%d %B %Y', '%d %B, %Y'):
                try:
                    return (time.strftime('%Y-%m-%d', time.strptime(m.group(1), fmt)),
                            'page text', versions)
                except ValueError:
                    pass
        return None, 'not found on page', versions

    def quit(self):
        if self.driver is not None:
            self.driver.quit()
            self.driver = None


# --------------------------------------------------------------------------- #
# spreadsheet plumbing
# --------------------------------------------------------------------------- #
def load_rows(path, sheet_name):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        sys.exit(f'Sheet {sheet_name!r} not found. Available: {wb.sheetnames}')
    ws = wb[sheet_name]
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else '' for c in next(it)]
    rows = [dict(zip(header, r)) for r in it if any(c is not None for c in r)]
    wb.close()
    return header, rows


def load_cache(path):
    if os.path.exists(path):
        with open(path, encoding='utf-8') as f:
            return json.load(f)
    return {}


def save_cache(path, cache):
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(cache, f, indent=1, sort_keys=True)
    os.replace(tmp, path)


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--input', default=DEFAULT_INPUT)
    ap.add_argument('--sheet', default=DEFAULT_SHEET)
    ap.add_argument('--output', default='V1_Publish_Dates.xlsx')
    ap.add_argument('--doi-col', default='DOI')
    ap.add_argument('--source-col', default='Source')
    ap.add_argument('--date-col', default='Published Date')
    ap.add_argument('--url-col', default='URL')
    ap.add_argument('--source', default='Research Square',
                    help="only process rows from this source; ANY disables the filter "
                         "(use on the 'Research Square' sheet, where some rows have a "
                         "blank Source cell)")
    ap.add_argument('--cache', default=CACHE_FILE)
    ap.add_argument('--limit', type=int, help='only process the first N DOIs (testing)')
    ap.add_argument('--sleep', type=float, default=0.2, help='delay between requests')
    ap.add_argument('--selenium', action='store_true',
                    help='allow the Chrome fallback for DOIs the APIs cannot resolve')
    ap.add_argument('--refresh', action='store_true', help='ignore the cache')
    args = ap.parse_args()

    header, rows = load_rows(args.input, args.sheet)
    for col in (args.doi_col, args.source_col, args.date_col):
        if col not in header:
            sys.exit(f'Column {col!r} not in sheet. Columns: {header}')

    # Research Square rows whose recorded DOI is a version other than v1.
    any_source = args.source.strip().upper() == "ANY"
    targets = []
    skipped_v1 = other_source = 0
    for i, row in enumerate(rows):
        doi = str(row.get(args.doi_col) or '').strip()
        if not doi:
            continue
        if (not any_source) and str(row.get(args.source_col) or '').strip() != args.source:
            other_source += 1
            continue
        v = version_of(doi)
        if v == 1:
            skipped_v1 += 1
        elif v is not None:
            targets.append((i, doi, row))

    print(f'{len(rows)} rows in {args.sheet!r}')
    if not any_source:
        print(f'  {other_source} rows from other sources          -> skipped')
    label = 'Research Square' if any_source else args.source
    print(f'  {skipped_v1} {label} DOIs already at /v1  -> skipped')
    print(f'  {len(targets)} {label} DOIs past /v1      -> fetching v1 date')

    if args.limit:
        targets = targets[:args.limit]

    cache = {} if args.refresh else load_cache(args.cache)
    session = make_session()
    browser = Browser()

    results = {}
    counts = {'crossref': 0, 'web': 0, 'selenium': 0, 'cached': 0, 'failed': 0}

    for n, (idx, doi, row) in enumerate(targets, 1):
        if doi in cache and cache[doi].get('v1_date'):
            results[idx] = cache[doi]
            counts['cached'] += 1
            continue

        v1_doi = to_v1(doi)
        sheet_url = row.get(args.url_col)
        rec = {'v1_doi': v1_doi, 'v1_url': rs_article_url(doi, 'v1', sheet_url),
               'v1_date': None, 'method': None, 'note': None, 'versions': {}}

        try:
            d, note = date_from_crossref(session, v1_doi)
            if d:
                rec.update(v1_date=d, method='crossref', note=note)
                counts['crossref'] += 1
            else:
                rec['note'] = f'crossref: {note}'
        except Exception as e:
            rec['note'] = f'crossref error: {e}'

        if not rec['v1_date']:
            try:
                d, note, versions = date_from_web(session, doi, sheet_url)
                rec['versions'] = {str(k): v for k, v in versions.items()}
                if d:
                    rec.update(v1_date=d, method='web', note=note)
                    counts['web'] += 1
                else:
                    rec['note'] = f'web: {note}'
            except Exception as e:
                rec['note'] = f'web error: {e}'

        if not rec['v1_date'] and args.selenium:
            try:
                d, note, versions = browser.date_for(doi, sheet_url)
                if versions:
                    rec['versions'] = {str(k): v for k, v in versions.items()}
                if d:
                    rec.update(v1_date=d, method='selenium', note=note)
                    counts['selenium'] += 1
                else:
                    rec['note'] = f'selenium: {note}'
            except Exception as e:
                rec['note'] = f'selenium error: {e}'

        if not rec['v1_date']:
            counts['failed'] += 1

        cache[doi] = rec
        results[idx] = rec

        if n % 25 == 0 or n == len(targets):
            save_cache(args.cache, cache)
            print(f'  [{n}/{len(targets)}] {counts}', flush=True)
        time.sleep(args.sleep)

    save_cache(args.cache, cache)
    browser.quit()

    # ---- write the output workbook -------------------------------------- #
    new_cols = ['V1 DOI', 'V1 URL', 'V1 Published Date', 'Recorded Version',
                'Days V1 to Recorded Version', 'Total Versions Seen',
                'V1 Date Source', 'Notes']
    out = openpyxl.Workbook()

    ws = out.active
    ws.title = 'All rows'
    ws.append(header + new_cols)

    hits = out.create_sheet('Non-v1 only')
    hits.append(header + new_cols)

    for i, row in enumerate(rows):
        base = [row.get(h) for h in header]
        rec = results.get(i)
        if not rec:
            ws.append(base + [''] * len(new_cols))
            continue
        v1_date = rec.get('v1_date')
        extra = [
            rec.get('v1_doi', ''),
            rec.get('v1_url', ''),
            v1_date or '',
            version_of(str(row.get(args.doi_col) or '')) or '',
            days_between(v1_date, row.get(args.date_col)) if v1_date else '',
            len(rec.get('versions') or {}) or '',
            rec.get('method') or '',
            rec.get('note') or '',
        ]
        ws.append(base + extra)
        hits.append(base + extra)

    for sheet in (ws, hits):
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions['B'].width = 34
        sheet.column_dimensions['C'].width = 55
    out.save(args.output)

    resolved = sum(1 for r in results.values() if r.get('v1_date'))
    print(f'\nResolved {resolved}/{len(targets)} v1 dates  {counts}')
    print(f'Wrote {args.output}  (sheets: "All rows", "Non-v1 only")')
    if counts['failed']:
        print('Unresolved:')
        for rec in results.values():
            if not rec.get('v1_date'):
                print(f'  {rec["v1_doi"]}  {rec.get("note")}')


if __name__ == '__main__':
    main()
